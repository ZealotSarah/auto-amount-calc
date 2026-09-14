"""基金金额自动测算工具。

结果 Sheet 仅输出用户约定的九列汇总字段。固定比例规则下住院和门诊
均使用内置的《24年25年目录内（费用）住院基金支付比例》。
"""

from __future__ import annotations

import datetime as dt
import os
import posixpath
import re
import shutil
import tempfile
import threading
import zipfile
import xml.etree.ElementTree as ET
from collections import OrderedDict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except ImportError:  # 无 GUI 环境（如无 Tk 的测试环境）下仍可导入核心计算逻辑。
    import types
    tk = types.SimpleNamespace(Tk=object)
    filedialog = messagebox = ttk = None


RATE_VERSION = "24-25 年住院基金支付比例"
APP_VERSION = "0.9.1"
OUTPUT_SHEET = "基金测算"
AUTO_SOURCE_SHEET = "自动选择（仅唯一匹配时）"
SUMMARY_HEADERS = ("医疗机构编码", "医疗机构名称", "险种类别", "医疗类别", "医疗总额", "数量总和", "人次", "基金金额", "报销比例")

SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOCUMENT_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

COLUMN_CANDIDATES: dict[str, tuple[str, ...]] = {
    "code": ("定点编码", "医疗机构编码"),
    "name": ("定点名称", "医疗机构名称"),
    "insurance": ("险种类型", "险种类别", "INSUTYPE", "insutype"),
    "medical_total": ("医疗费总额", "医疗总额"),
    "fund_total": ("基金支付总额",),
    "scope_amount": ("符合范围金额",),
    "visit_type": ("医疗类别",),
    "visit_name": ("医疗类别名称",),
    "violation_amount": ("违规金额",),
    "price": ("单价",),
    "quantity": ("数量",),
}
REQUIRED_COLUMNS = ("code", "insurance", "medical_total", "fund_total", "scope_amount", "quantity")
COLUMN_LABELS = {
    "code": "定点编码", "name": "定点名称", "insurance": "险种类型",
    "medical_total": "医疗费总额", "fund_total": "基金支付总额",
    "scope_amount": "符合范围金额", "visit_type": "医疗类别",
    "visit_name": "医疗类别名称", "violation_amount": "违规金额",
    "price": "单价", "quantity": "数量",
}

# 统筹区: {医疗机构级别: (职工, 居民, 合计)}。数值为小数比例。
FUND_RATES: dict[str, dict[str, tuple[Decimal, Decimal, Decimal]]] = {
    "河北省": {"三级": (Decimal("0.8423"), Decimal("0.6112"), Decimal("0.6860")), "二级": (Decimal("0.8882"), Decimal("0.7227"), Decimal("0.7550")), "一级": (Decimal("0.9109"), Decimal("0.8422"), Decimal("0.8511"))},
    "石家庄市": {"三级": (Decimal("0.8407"), Decimal("0.6158"), Decimal("0.6946")), "二级": (Decimal("0.9169"), Decimal("0.7657"), Decimal("0.7958")), "一级": (Decimal("0.9157"), Decimal("0.8730"), Decimal("0.8796"))},
    "辛集市": {"三级": (Decimal("0.8016"), Decimal("0.4953"), Decimal("0.5507")), "二级": (Decimal("0.8278"), Decimal("0.6711"), Decimal("0.6905")), "一级": (Decimal("0.8944"), Decimal("0.7424"), Decimal("0.7672"))},
    "唐山市": {"三级": (Decimal("0.8226"), Decimal("0.5689"), Decimal("0.6874")), "二级": (Decimal("0.8868"), Decimal("0.6471"), Decimal("0.7275")), "一级": (Decimal("0.9126"), Decimal("0.8150"), Decimal("0.8435"))},
    "秦皇岛市": {"三级": (Decimal("0.8368"), Decimal("0.5892"), Decimal("0.6892")), "二级": (Decimal("0.8974"), Decimal("0.7384"), Decimal("0.7970")), "一级": (Decimal("0.9029"), Decimal("0.8865"), Decimal("0.8899"))},
    "邯郸市": {"三级": (Decimal("0.8757"), Decimal("0.6893"), Decimal("0.7425")), "二级": (Decimal("0.8941"), Decimal("0.7667"), Decimal("0.7850")), "一级": (Decimal("0.9039"), Decimal("0.8510"), Decimal("0.8550"))},
    "邢台市": {"三级": (Decimal("0.8960"), Decimal("0.6712"), Decimal("0.7260")), "二级": (Decimal("0.9468"), Decimal("0.7669"), Decimal("0.7931")), "一级": (Decimal("0.9576"), Decimal("0.8614"), Decimal("0.8666"))},
    "保定市": {"三级": (Decimal("0.8238"), Decimal("0.5594"), Decimal("0.6386")), "二级": (Decimal("0.8496"), Decimal("0.7242"), Decimal("0.7406")), "一级": (Decimal("0.8874"), Decimal("0.8683"), Decimal("0.8703"))},
    "定州市": {"三级": (Decimal("0.8033"), Decimal("0.5159"), Decimal("0.5675")), "二级": (Decimal("0.8754"), Decimal("0.6263"), Decimal("0.6586")), "一级": (Decimal("0.9372"), Decimal("0.8123"), Decimal("0.8176"))},
    "张家口市": {"三级": (Decimal("0.8649"), Decimal("0.7032"), Decimal("0.7563")), "二级": (Decimal("0.8767"), Decimal("0.7908"), Decimal("0.8097")), "一级": (Decimal("0.9266"), Decimal("0.9030"), Decimal("0.9070"))},
    "承德市": {"三级": (Decimal("0.8555"), Decimal("0.5894"), Decimal("0.6625")), "二级": (Decimal("0.8593"), Decimal("0.7169"), Decimal("0.7382")), "一级": (Decimal("0.8777"), Decimal("0.8701"), Decimal("0.8705"))},
    "沧州市": {"三级": (Decimal("0.8475"), Decimal("0.5926"), Decimal("0.6521")), "二级": (Decimal("0.8639"), Decimal("0.7189"), Decimal("0.7388")), "一级": (Decimal("0.9258"), Decimal("0.8626"), Decimal("0.8672"))},
    "廊坊市": {"三级": (Decimal("0.8458"), Decimal("0.5684"), Decimal("0.6424")), "二级": (Decimal("0.9131"), Decimal("0.6778"), Decimal("0.7155")), "一级": (Decimal("0.9581"), Decimal("0.8326"), Decimal("0.8484"))},
    "衡水市": {"三级": (Decimal("0.8063"), Decimal("0.5892"), Decimal("0.6346")), "二级": (Decimal("0.8331"), Decimal("0.6632"), Decimal("0.6866")), "一级": (Decimal("0.8830"), Decimal("0.7769"), Decimal("0.7854"))},
    "雄安新区": {"三级": (Decimal("0.8283"), Decimal("0.6516"), Decimal("0.6719")), "二级": (Decimal("0.8589"), Decimal("0.6552"), Decimal("0.6713")), "一级": (Decimal("0.8462"), Decimal("0.6528"), Decimal("0.6537"))},
    "省本级": {"三级": (Decimal("0.8649"), Decimal("0.8649"), Decimal("0.8649")), "二级": (Decimal("0.9150"), Decimal("0.9150"), Decimal("0.9150")), "一级": (Decimal("0.8975"), Decimal("0.8975"), Decimal("0.8975"))},
}


class CalculationError(Exception):
    """A business-data error that should be shown to the operator."""


@dataclass
class RunOptions:
    rule_type: str
    visit_type: str
    pooling_area: str
    institution_level: str
    deduction_price: Decimal | None
    deduction_quantity: Decimal | None
    overwrite_result: bool
    source_sheet: str | None = None


@dataclass
class FileResult:
    path: Path
    sheet_name: str
    processed: int
    successful: int
    errors: int
    excluded: int
    total_fund: Decimal
    warnings: list[str]
    backup_path: Path | None = None


def clean_header(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).replace("　", "")


def as_decimal(value: Any, label: str) -> Decimal:
    if value is None or str(value).strip() == "":
        raise CalculationError(f"{label}为空")
    try:
        number = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError) as exc:
        raise CalculationError(f"{label}不是数字：{value}") from exc
    if not number.is_finite():
        raise CalculationError(f"{label}不是有限数字：{value}")
    return number


def optional_decimal(value: str, label: str) -> Decimal | None:
    if not value.strip():
        return None
    return as_decimal(value, label)


def text_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def collect_headers(cells: Iterable[Any]) -> dict[str, list[int]]:
    headers: dict[str, list[int]] = {}
    for cell in cells:
        header = clean_header(cell.value)
        if header:
            headers.setdefault(header, []).append(cell.column)
    return headers


def find_column(headers: dict[str, list[int]], candidates: Iterable[str], label: str = "字段") -> int | None:
    exact_matches: list[tuple[str, int]] = []
    for candidate in candidates:
        exact = clean_header(candidate)
        exact_matches.extend((exact, index) for index in headers.get(exact, []))
    if exact_matches:
        if len(exact_matches) > 1:
            locations = "、".join(f"{name}（第 {index} 列）" for name, index in exact_matches)
            raise CalculationError(f"{label}匹配到多个列：{locations}")
        return exact_matches[0][1]

    prefix_matches: list[tuple[str, int]] = []
    prefixes = tuple(clean_header(candidate) for candidate in candidates)
    for header, indexes in headers.items():
        if any(header.startswith(prefix) for prefix in prefixes):
            prefix_matches.extend((header, index) for index in indexes)
    if len(prefix_matches) > 1:
        locations = "、".join(f"{name}（第 {index} 列）" for name, index in prefix_matches)
        raise CalculationError(f"{label}匹配到多个列：{locations}")
    return prefix_matches[0][1] if prefix_matches else None


def resolve_columns(headers: dict[str, list[int]]) -> dict[str, int | None]:
    return {
        name: find_column(headers, candidates, COLUMN_LABELS[name])
        for name, candidates in COLUMN_CANDIDATES.items()
    }


def determine_source_sheet(workbook, requested_sheet: str | None = None, required_columns: tuple[str, ...] = REQUIRED_COLUMNS) -> tuple[Any, int, dict[str, list[int]]]:
    if requested_sheet:
        if requested_sheet not in workbook.sheetnames:
            raise CalculationError(f"找不到指定的源数据 Sheet：{requested_sheet}")
        sheets = [workbook[requested_sheet]]
    else:
        sheets = [sheet for sheet in workbook.worksheets if not sheet.title.startswith(OUTPUT_SHEET)]

    valid: list[tuple[Any, int, dict[str, list[int]]]] = []
    diagnostics: list[str] = []
    for sheet in sheets:
        best: tuple[int, int, dict[str, list[int]], list[str]] | None = None
        for header_row in range(1, min(10, sheet.max_row) + 1):
            headers = collect_headers(sheet[header_row])
            try:
                columns = resolve_columns(headers)
            except CalculationError as exc:
                diagnostics.append(f"{sheet.title} 第 {header_row} 行：{exc}")
                continue
            missing = [COLUMN_LABELS[name] for name in required_columns if columns[name] is None]
            score = len(required_columns) - len(missing)
            candidate = (score, -header_row, headers, missing)
            if best is None or candidate[:2] > best[:2]:
                best = candidate
        if best is None:
            continue
        if not best[3]:
            valid.append((sheet, -best[1], best[2]))
        else:
            diagnostics.append(f"{sheet.title}：缺少" + "、".join(best[3]))

    if len(valid) == 1:
        return valid[0]
    if len(valid) > 1:
        names = "、".join(sheet.title for sheet, _, _ in valid)
        raise CalculationError(f"检测到多个可用业务 Sheet：{names}。请在界面中明确选择源数据 Sheet。")
    detail = "；".join(diagnostics[:5])
    message = "找不到字段完整且唯一的业务数据 Sheet。"
    raise CalculationError(message + (f" {detail}" if detail else ""))


def insurance_bucket(insurance: Any) -> int | None:
    text = text_value(insurance)
    if text in {"310", "职工", "城镇职工", "职工基本医疗保险", "城镇职工基本医疗保险"}:
        return 0
    if text in {"390", "城乡", "居民", "城乡居民", "城乡居民基本医疗保险"}:
        return 1
    return None


def resolved_visit_type(raw: Any, selected: str, name: Any = None) -> str:
    if selected != "自动识别":
        return selected
    value = text_value(raw)
    code_types = {"11": "门诊", "14": "门诊", "51": "门诊", "21": "住院", "22": "住院", "52": "住院"}
    if value in code_types:
        return code_types[value]
    for label in (value, text_value(name)):
        if "住院" in label:
            return "住院"
        if "门诊" in label:
            return "门诊"
    raise CalculationError(f"无法从医疗类别识别住院或门诊：{value or '空值'}")


def validate_options(options: RunOptions) -> None:
    if options.rule_type not in {"通用", "串换", "固定比例"}:
        raise CalculationError("请选择规则大类。")
    if options.visit_type not in {"自动识别", "住院", "门诊"}:
        raise CalculationError("请选择业务类型。")
    if options.pooling_area not in FUND_RATES:
        raise CalculationError("请选择参保地（比例表匹配）。")
    if options.institution_level not in {"三级", "二级", "一级"}:
        raise CalculationError("请选择医疗机构级别。")
    for value, label in ((options.deduction_price, "扣减单价"), (options.deduction_quantity, "扣减数量")):
        if value is not None and (not value.is_finite() or value < 0):
            raise CalculationError(f"{label}必须是大于或等于 0 的有限数字。")
    if options.rule_type == "串换" and options.deduction_quantity is None:
        raise CalculationError("串换规则必须填写扣减数量。")


def worksheet_parts(package: dict[str, bytes]) -> dict[str, str]:
    workbook_root = ET.fromstring(package["xl/workbook.xml"])
    relationships_root = ET.fromstring(package["xl/_rels/workbook.xml.rels"])
    relationship_targets = {
        relationship.attrib["Id"]: relationship.attrib["Target"]
        for relationship in relationships_root.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
    }
    parts: dict[str, str] = {}
    for sheet in workbook_root.findall(f".//{{{SPREADSHEET_NS}}}sheet"):
        relationship_id = sheet.attrib[f"{{{DOCUMENT_REL_NS}}}id"]
        target = relationship_targets[relationship_id]
        part = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
        parts[sheet.attrib["name"]] = part
    return parts


def capture_formula_caches(path: Path) -> dict[str, dict[str, tuple[str, str | None]]]:
    with zipfile.ZipFile(path) as archive:
        package = {name: archive.read(name) for name in archive.namelist()}
    caches: dict[str, dict[str, tuple[str, str | None]]] = {}
    for sheet_name, part in worksheet_parts(package).items():
        root = ET.fromstring(package[part])
        sheet_caches: dict[str, tuple[str, str | None]] = {}
        for cell in root.findall(f".//{{{SPREADSHEET_NS}}}c"):
            formula = cell.find(f"{{{SPREADSHEET_NS}}}f")
            value = cell.find(f"{{{SPREADSHEET_NS}}}v")
            if formula is not None and value is not None and value.text is not None:
                sheet_caches[cell.attrib["r"]] = (value.text, cell.attrib.get("t"))
        if sheet_caches:
            caches[sheet_name] = sheet_caches
    return caches


def restore_formula_caches(path: Path, caches: dict[str, dict[str, tuple[str, str | None]]]) -> None:
    if not caches:
        return
    with zipfile.ZipFile(path) as archive:
        entries = [(item, archive.read(item.filename)) for item in archive.infolist()]
    package = {item.filename: data for item, data in entries}
    changed_parts: dict[str, bytes] = {}
    for sheet_name, part in worksheet_parts(package).items():
        sheet_caches = caches.get(sheet_name)
        if not sheet_caches:
            continue
        root = ET.fromstring(package[part])
        changed = False
        for cell in root.findall(f".//{{{SPREADSHEET_NS}}}c"):
            cached = sheet_caches.get(cell.attrib.get("r", ""))
            formula = cell.find(f"{{{SPREADSHEET_NS}}}f")
            if cached is None or formula is None:
                continue
            value = cell.find(f"{{{SPREADSHEET_NS}}}v")
            if value is None:
                value = ET.Element(f"{{{SPREADSHEET_NS}}}v")
                children = list(cell)
                cell.insert(children.index(formula) + 1, value)
            value.text, cell_type = cached
            if cell_type is None:
                cell.attrib.pop("t", None)
            else:
                cell.attrib["t"] = cell_type
            changed = True
        if changed:
            changed_parts[part] = ET.tostring(root, encoding="utf-8", xml_declaration=False)
    if not changed_parts:
        return

    descriptor, rewritten_name = tempfile.mkstemp(prefix=f".{path.stem}_formula_", suffix=path.suffix, dir=path.parent)
    os.close(descriptor)
    rewritten_path = Path(rewritten_name)
    try:
        with zipfile.ZipFile(rewritten_path, "w") as archive:
            for item, data in entries:
                archive.writestr(item, changed_parts.get(item.filename, data))
        os.replace(rewritten_path, path)
    finally:
        if rewritten_path.exists():
            rewritten_path.unlink()


def run_file(path: Path, options: RunOptions) -> FileResult:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise CalculationError("仅支持 .xlsx 和 .xlsm 文件。")
    validate_options(options)

    keep_vba = path.suffix.lower() == ".xlsm"
    formula_caches = capture_formula_caches(path)
    workbook = load_workbook(path, keep_vba=keep_vba)
    try:
        return calculate_workbook(workbook, path, options, keep_vba, formula_caches)
    finally:
        workbook.close()


def calculate_workbook(workbook, path: Path, options: RunOptions, keep_vba: bool, formula_caches: dict[str, dict[str, tuple[str, str | None]]] | None = None) -> FileResult:
    required_columns = tuple(
        name for name in REQUIRED_COLUMNS
        if not (options.rule_type == "固定比例" and name == "fund_total")
    )
    source, header_row, headers = determine_source_sheet(workbook, options.source_sheet, required_columns)
    columns = resolve_columns(headers)
    missing = [name for name in required_columns if columns[name] is None]
    if missing:
        raise CalculationError("缺少必填列：" + "、".join(COLUMN_LABELS[name] for name in missing))
    if options.visit_type == "自动识别" and columns["visit_type"] is None:
        raise CalculationError("选择“自动识别”时，原表必须有“医疗类别”列。")
    if options.rule_type == "串换" and columns["violation_amount"] is None:
        if columns["price"] is None:
            raise CalculationError("串换规则且无“违规金额”列时，源数据必须有“单价”列。")
        if options.deduction_price is None:
            raise CalculationError("串换规则且无“违规金额”列时，必须填写扣减单价。")

    numeric_columns = [columns[name] for name in ("medical_total", "fund_total", "scope_amount", "quantity", "violation_amount", "price") if columns[name]]
    has_numeric_formulas = any(
        source.cell(row, column).data_type == "f"
        for row in range(header_row + 1, source.max_row + 1)
        for column in numeric_columns
    )
    values_workbook = load_workbook(path, data_only=True, keep_vba=keep_vba) if has_numeric_formulas else None
    value_source = values_workbook[source.title] if values_workbook else source

    output_name = create_output_sheet_name(workbook, options.overwrite_result)
    if output_name == OUTPUT_SHEET and OUTPUT_SHEET in workbook.sheetnames:
        del workbook[OUTPUT_SHEET]
    output = workbook.create_sheet(output_name)
    output_name = output.title

    groups: OrderedDict[tuple[str, str, str], dict[str, Any]] = OrderedDict()
    institution_names: dict[str, str] = {}
    warnings: list[str] = []
    processed = successful = errors = excluded = 0
    formula_rows = source.iter_rows(min_row=header_row + 1)
    value_rows = value_source.iter_rows(min_row=header_row + 1, values_only=True)
    for source_row, (formula_row, row) in enumerate(zip(formula_rows, value_rows), header_row + 1):
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        processed += 1
        get = lambda name: row[columns[name] - 1] if columns[name] else None
        get_cell = lambda name: formula_row[columns[name] - 1] if columns[name] else None

        def get_decimal(name: str, label: str) -> Decimal:
            value = get(name)
            cell = get_cell(name)
            if cell is not None and cell.data_type == "f" and (value is None or str(value).strip() == ""):
                raise CalculationError(f"{label}公式没有缓存结果，请先用 Excel 打开并保存源文件")
            return as_decimal(value, label)

        code = text_value(get("code")) or "未填写编码"
        name = text_value(get("name"))
        if name:
            institution_names.setdefault(code, name)
        insurance = text_value(get("insurance")) or "未填写险种"
        bucket = insurance_bucket(insurance)
        if bucket is None:
            excluded += 1
            continue
        try:
            medical_total = get_decimal("medical_total", "医疗费总额")
            fund_total = Decimal("0") if options.rule_type == "固定比例" else get_decimal("fund_total", "基金支付总额")
            scope_amount = get_decimal("scope_amount", "符合范围金额")
            quantity = get_decimal("quantity", "数量")
            if medical_total < 0 or fund_total < 0 or scope_amount < 0 or quantity < 0:
                raise CalculationError("医疗费总额、基金支付总额、符合范围金额和数量不能为负数")
            visit_type = resolved_visit_type(get("visit_type"), options.visit_type, get("visit_name"))
            result_quantity = quantity
            if options.rule_type == "串换":
                assert options.deduction_quantity is not None
                result_quantity -= options.deduction_quantity
                if result_quantity < 0:
                    raise CalculationError("串换扣减后的数量为负数")

            if options.rule_type == "固定比例":
                base = scope_amount
            elif visit_type == "门诊":
                if medical_total == 0:
                    raise CalculationError("门诊医疗费总额为 0，无法计算")
                base = scope_amount
            else:
                if options.rule_type == "通用":
                    base = scope_amount
                elif columns["violation_amount"] is not None:
                    base = get_decimal("violation_amount", "违规金额")
                else:
                    price = get_decimal("price", "原单价")
                    assert options.deduction_price is not None
                    result_price = price - options.deduction_price
                    if result_price < 0:
                        raise CalculationError("串换扣减后的单价为负数")
                    base = result_price * result_quantity
                if base < 0:
                    raise CalculationError("计算基数为负数")

            key = (code, insurance, visit_type)
            group = groups.setdefault(key, {
                "name": "未填写名称", "medical_total": Decimal("0"), "fund_total": Decimal("0"),
                "calculation_base": Decimal("0"), "quantity_total": Decimal("0"),
                "count": 0,
                "fund_amount": Decimal("0"),
                "types": {visit_type}, "rates": set(),
            })
            group["medical_total"] += medical_total
            group["fund_total"] += fund_total
            group["calculation_base"] += base
            group["quantity_total"] += result_quantity
            group["count"] += 1
            successful += 1
        except CalculationError as exc:
            errors += 1
            if len(warnings) < 30:
                warnings.append(f"第 {source_row} 行：{exc}")

    if values_workbook:
        values_workbook.close()
    if successful == 0 and errors > 0:
        detail = warnings[0] if warnings else "所有明细均计算失败"
        raise CalculationError(f"没有成功计算的明细：{detail}")

    for (code, _, _), group in groups.items():
        group["name"] = institution_names.get(code, "未填写名称")

    total_fund = Decimal("0")
    total_count = 0
    for (_, insurance, visit_type), group in groups.items():
        bucket = insurance_bucket(insurance)
        assert bucket is not None
        if options.rule_type == "固定比例":
            rate = FUND_RATES[options.pooling_area][options.institution_level][bucket]
        elif visit_type == "门诊":
            rate = group["fund_total"] / group["medical_total"]
        else:
            rate = FUND_RATES[options.pooling_area][options.institution_level][bucket]
        amount = group["calculation_base"] * rate
        group["fund_amount"] = amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        group["rates"].add(rate)
        total_fund += group["fund_amount"]
        total_count += group["count"]
    if total_count != successful:
        raise CalculationError(f"人次勾稽校验失败：分组人次合计 {total_count} 与计算成功数 {successful} 不一致。")

    write_output_sheet(output, path, source.title, options, groups, processed, successful, errors, total_fund, warnings, excluded)
    backup_path = save_workbook_safely(
        workbook, path, output_name, total_fund, successful, len(groups), keep_vba,
        formula_caches or {},
    )
    return FileResult(path, output_name, processed, successful, errors, excluded, total_fund, warnings, backup_path)


def create_output_sheet_name(workbook, overwrite: bool) -> str:
    if OUTPUT_SHEET not in workbook.sheetnames or overwrite:
        return OUTPUT_SHEET
    suffix = dt.datetime.now().strftime("%m%d_%H%M%S")
    return f"{OUTPUT_SHEET}_{suffix}"


def next_backup_path(path: Path) -> Path:
    timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    base = path.with_name(f"{path.stem}_测算前备份_{timestamp}{path.suffix}")
    candidate = base
    number = 2
    while candidate.exists():
        candidate = path.with_name(f"{path.stem}_测算前备份_{timestamp}_{number}{path.suffix}")
        number += 1
    return candidate


def replace_with_backup(path: Path, temp_path: Path, backup_path: Path) -> None:
    if os.name == "nt":
        import ctypes

        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        replace_file = kernel32.ReplaceFileW
        replace_file.argtypes = [
            ctypes.c_wchar_p, ctypes.c_wchar_p, ctypes.c_wchar_p,
            ctypes.c_ulong, ctypes.c_void_p, ctypes.c_void_p,
        ]
        replace_file.restype = ctypes.c_int
        if not replace_file(str(path), str(temp_path), str(backup_path), 1, None, None):
            error = ctypes.get_last_error()
            raise OSError(error, ctypes.FormatError(error), str(path))
        return

    shutil.copy2(path, backup_path)
    os.replace(temp_path, path)


def labeled_value(sheet, label: str, max_row: int = 6, max_column: int = 11) -> Any:
    for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column):
        for cell in row:
            if cell.value == label:
                return sheet.cell(cell.row, cell.column + 1).value
    raise CalculationError(f"保存校验失败：找不到“{label}”，原文件未替换。")


def save_workbook_safely(
    workbook, path: Path, output_name: str, total_fund: Decimal, successful: int,
    group_count: int, keep_vba: bool,
    formula_caches: dict[str, dict[str, tuple[str, str | None]]],
) -> Path:
    descriptor, temp_name = tempfile.mkstemp(prefix=f".{path.stem}_", suffix=path.suffix, dir=path.parent)
    os.close(descriptor)
    temp_path = Path(temp_name)
    backup_path: Path | None = None
    try:
        workbook.save(temp_path)
        workbook.close()
        restore_formula_caches(temp_path, formula_caches)

        verification = load_workbook(temp_path, read_only=True, data_only=True, keep_vba=keep_vba)
        try:
            if output_name not in verification.sheetnames:
                raise CalculationError("保存校验失败：结果 Sheet 缺失，原文件未替换。")
            saved_output = verification[output_name]
            saved_headers = tuple(saved_output.cell(7, column).value for column in range(1, len(SUMMARY_HEADERS) + 1))
            if saved_headers != SUMMARY_HEADERS:
                raise CalculationError("保存校验失败：九列汇总表头不一致，原文件未替换。")
            saved_total = as_decimal(labeled_value(saved_output, "基金金额合计"), "保存后的基金金额合计")
            if saved_total.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) != total_fund:
                raise CalculationError("保存校验失败：基金金额合计不一致，原文件未替换。")
            saved_count = sum(
                int(as_decimal(saved_output.cell(row, 7).value, "保存后的人次"))
                for row in range(8, 8 + group_count)
            )
            if saved_count != successful:
                raise CalculationError(
                    f"保存校验失败：人次合计 {saved_count} 与计算成功数 {successful} 不一致，原文件未替换。"
                )
        finally:
            verification.close()

        backup_path = next_backup_path(path)
        replace_with_backup(path, temp_path, backup_path)
        return backup_path
    except PermissionError as exc:
        raise CalculationError(f"无法替换原文件。请关闭 Excel 中已打开的工作簿后重试：{path.name}") from exc
    except OSError as exc:
        raise CalculationError(f"安全保存失败，原文件未替换：{exc}") from exc
    finally:
        workbook.close()
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass


def write_output_sheet(output, path: Path, source_name: str, options: RunOptions, groups: OrderedDict, processed: int, successful: int, errors: int, total_fund: Decimal, warnings: list[str], excluded: int) -> None:
    navy = PatternFill("solid", fgColor="1F4E78")
    blue = PatternFill("solid", fgColor="D9EAF7")
    output.merge_cells("A1:I1")
    output["A1"] = "基金金额测算结果"
    output["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    output["A1"].fill = navy
    output["A1"].alignment = Alignment(horizontal="center")
    metadata = [
        ("源文件", path.name), ("源数据 Sheet", source_name), ("规则大类", options.rule_type),
        ("业务类型", options.visit_type), ("参保地（比例表匹配）", options.pooling_area),
        ("医疗机构级别", options.institution_level), ("比例版本", RATE_VERSION),
        ("处理记录数", processed), ("计算成功数", successful), ("异常数", errors),
        ("基金金额合计", total_fund),
        ("排除险种记录数", excluded),
    ]
    for index, (label, value) in enumerate(metadata):
        row = 2 + index // 3
        column = 1 + (index % 3) * 2
        output.cell(row, column, label).font = Font(bold=True)
        output.cell(row, column).fill = blue
        output.cell(row, column + 1, value)
    output["D5"].number_format = "#,##0.00"

    audit_metadata = [
        ("生成时间", dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("工具版本", APP_VERSION),
        ("扣减单价", options.deduction_price if options.deduction_price is not None else "未填写"),
        ("扣减数量", options.deduction_quantity if options.deduction_quantity is not None else "未填写"),
        ("输出形式", "普通汇总表（非 Excel 原生透视表）"),
    ]
    for row, (label, value) in enumerate(audit_metadata, 2):
        output.cell(row, 10, label).font = Font(bold=True)
        output.cell(row, 10).fill = blue
        output.cell(row, 11, value)

    header_row = 7
    for column, header in enumerate(SUMMARY_HEADERS, 1):
        cell = output.cell(header_row, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = navy
        cell.alignment = Alignment(horizontal="center")
    for row_number, ((code, insurance, visit_type), group) in enumerate(groups.items(), header_row + 1):
        rate = group_rate(group)
        output.cell(row_number, 1, code)
        output.cell(row_number, 2, group["name"])
        output.cell(row_number, 3, insurance)
        output.cell(row_number, 4, visit_type)
        output.cell(row_number, 5, float(group["medical_total"]))
        output.cell(row_number, 6, float(group["quantity_total"]))
        output.cell(row_number, 7, group["count"])
        output.cell(row_number, 8, float(group["fund_amount"]))
        output.cell(row_number, 9, float(rate) if isinstance(rate, Decimal) else rate)
        output.cell(row_number, 5).number_format = "#,##0.00"
        output.cell(row_number, 6).number_format = "#,##0.####"
        output.cell(row_number, 7).number_format = "#,##0"
        output.cell(row_number, 8).number_format = "#,##0.00"
        if isinstance(rate, Decimal):
            output.cell(row_number, 9).number_format = "0.00%"
    if groups:
        last_row = header_row + len(groups)
        table = Table(displayName=unique_table_name(output), ref=f"A{header_row}:I{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        output.add_table(table)
    else:
        output.cell(header_row + 1, 1, "没有可汇总的数据")

    output.column_dimensions["A"].width = 20
    output.column_dimensions["B"].width = 46
    output.column_dimensions["C"].width = 28
    output.column_dimensions["D"].width = 14
    output.column_dimensions["E"].width = 16
    output.column_dimensions["F"].width = 16
    output.column_dimensions["G"].width = 12
    output.column_dimensions["H"].width = 16
    output.column_dimensions["I"].width = 14
    output.column_dimensions["J"].width = 18
    output.column_dimensions["K"].width = 34
    output.freeze_panes = "A8"
    if warnings:
        output["A" + str(header_row + len(groups) + 3)] = "异常示例（最多显示 30 条）"
        output["A" + str(header_row + len(groups) + 3)].font = Font(bold=True, color="C00000")
        for offset, warning in enumerate(warnings, 4):
            output["A" + str(header_row + len(groups) + offset)] = warning
            output.merge_cells(start_row=header_row + len(groups) + offset, start_column=1, end_row=header_row + len(groups) + offset, end_column=9)


def group_rate(group: dict[str, Any]) -> Decimal | str:
    types = group["types"]
    if len(group["rates"]) == 1:
        return next(iter(group["rates"]))
    if types == {"门诊"}:
        return Decimal("0") if group["medical_total"] == 0 else group["fund_total"] / group["medical_total"]
    return "混合类型"


def unique_table_name(sheet) -> str:
    """Excel table names are workbook-wide, even on different sheets."""
    existing = {name for worksheet in sheet.parent.worksheets for name in worksheet.tables}
    base = "FundCalculationSummary"
    if base not in existing:
        return base
    number = 2
    while f"{base}{number}" in existing:
        number += 1
    return f"{base}{number}"


class CalculatorApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title(f"基金金额自动测算工具 v{APP_VERSION}")
        self.geometry("760x610")
        self.minsize(680, 560)
        self.files: list[Path] = []
        self.running = False
        self.rule_type = tk.StringVar(value="通用")
        self.visit_type = tk.StringVar(value="自动识别")
        self.pooling_area = tk.StringVar()
        self.institution_level = tk.StringVar()
        self.source_sheet = tk.StringVar(value=AUTO_SOURCE_SHEET)
        self.deduction_price = tk.StringVar()
        self.deduction_quantity = tk.StringVar()
        self.overwrite_result = tk.BooleanVar(value=False)
        self._build()
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    def _build(self) -> None:
        container = ttk.Frame(self, padding=16)
        container.pack(fill="both", expand=True)
        ttk.Label(container, text="基金金额自动测算工具", font=("Microsoft YaHei UI", 16, "bold")).pack(anchor="w")
        ttk.Label(container, text="先按结果维度汇总；固定比例下住院、门诊均使用比例表，其余规则门诊动态计算。", foreground="#555555").pack(anchor="w", pady=(4, 12))
        file_bar = ttk.Frame(container)
        file_bar.pack(fill="x")
        self.pick_button = ttk.Button(file_bar, text="选择 Excel 文件", command=self.pick_files)
        self.pick_button.pack(side="left")
        self.clear_button = ttk.Button(file_bar, text="清空", command=self.clear_files)
        self.clear_button.pack(side="left", padx=8)
        self.file_label = ttk.Label(file_bar, text="尚未选择文件")
        self.file_label.pack(side="left", padx=8)

        params = ttk.LabelFrame(container, text="参数", padding=12)
        params.pack(fill="x", pady=12)
        self.add_combo(params, "规则大类", self.rule_type, ["通用", "串换", "固定比例"], 0, 0)
        self.add_combo(params, "业务类型", self.visit_type, ["自动识别", "住院", "门诊"], 0, 1)
        self.add_combo(params, "参保地（比例表匹配）", self.pooling_area, list(FUND_RATES), 1, 0)
        self.add_combo(params, "医疗机构级别", self.institution_level, ["三级", "二级", "一级"], 1, 1)
        ttk.Label(params, text="源数据 Sheet").grid(row=2, column=0, sticky="w", pady=(10, 0))
        self.source_sheet_combo = ttk.Combobox(params, textvariable=self.source_sheet, values=[AUTO_SOURCE_SHEET], width=26)
        self.source_sheet_combo.grid(row=2, column=1, sticky="ew", pady=(10, 0))
        ttk.Label(params, text="扣减单价（仅串换无违规金额列时）").grid(row=3, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(params, textvariable=self.deduction_price, width=26).grid(row=3, column=1, sticky="ew", pady=(8, 0))
        ttk.Label(params, text="扣减数量（串换必填）").grid(row=4, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(params, textvariable=self.deduction_quantity, width=26).grid(row=4, column=1, sticky="ew", pady=(8, 0))
        ttk.Checkbutton(params, text="覆盖已有“基金测算”Sheet（否则自动新建带时间的 Sheet）", variable=self.overwrite_result).grid(row=5, column=0, columnspan=2, sticky="w", pady=(10, 0))
        params.columnconfigure(1, weight=1)

        controls = ttk.Frame(container)
        controls.pack(fill="x")
        self.run_button = ttk.Button(controls, text="开始测算", command=self.start)
        self.run_button.pack(side="left")
        ttk.Label(controls, text="  输出写回所选原 Excel 文件，请先关闭已打开的文件。", foreground="#A33").pack(side="left")
        self.log = tk.Text(container, height=13, wrap="word", state="disabled", font=("Consolas", 10))
        self.log.pack(fill="both", expand=True, pady=(12, 0))

    def add_combo(self, parent, label: str, variable: tk.StringVar, values: list[str], row: int, column: int) -> None:
        base = column * 2
        ttk.Label(parent, text=label).grid(row=row, column=base, sticky="w", padx=(0, 10), pady=4)
        combo = ttk.Combobox(parent, textvariable=variable, values=values, state="readonly", width=26)
        combo.grid(row=row, column=base + 1, sticky="ew", padx=(0, 20), pady=4)
        parent.columnconfigure(base + 1, weight=1)

    def pick_files(self) -> None:
        selections = filedialog.askopenfilenames(title="选择需要测算的 Excel 文件", filetypes=[("Excel 文件", "*.xlsx *.xlsm")])
        self.files = [Path(item) for item in selections]
        self.file_label.configure(text=f"已选择 {len(self.files)} 个文件" if self.files else "尚未选择文件")
        choices = [AUTO_SOURCE_SHEET]
        if self.files:
            first = self.files[0]
            try:
                workbook = load_workbook(first, read_only=True, keep_vba=first.suffix.lower() == ".xlsm")
                try:
                    choices.extend(sheet.title for sheet in workbook.worksheets if not sheet.title.startswith(OUTPUT_SHEET))
                finally:
                    workbook.close()
            except Exception as exc:
                messagebox.showerror("读取失败", f"无法读取第一个文件的 Sheet 列表：{exc}")
        self.source_sheet_combo.configure(values=choices)
        self.source_sheet.set(AUTO_SOURCE_SHEET)

    def clear_files(self) -> None:
        self.files = []
        self.file_label.configure(text="尚未选择文件")

    def start(self) -> None:
        if not self.files:
            messagebox.showwarning("未选择文件", "请先选择至少一个 Excel 文件。")
            return
        try:
            options = RunOptions(
                self.rule_type.get(), self.visit_type.get(), self.pooling_area.get(), self.institution_level.get(),
                optional_decimal(self.deduction_price.get(), "扣减单价"),
                optional_decimal(self.deduction_quantity.get(), "扣减数量"), self.overwrite_result.get(),
                None if self.source_sheet.get() in {"", AUTO_SOURCE_SHEET} else self.source_sheet.get(),
            )
            validate_options(options)
        except CalculationError as exc:
            messagebox.showerror("参数错误", str(exc))
            return
        self.running = True
        self.run_button.configure(state="disabled")
        self.pick_button.configure(state="disabled")
        self.clear_button.configure(state="disabled")
        self.write_log("开始处理...\n")
        threading.Thread(target=self.worker, args=(tuple(self.files), options), daemon=True).start()

    def worker(self, files: tuple[Path, ...], options: RunOptions) -> None:
        results: list[FileResult] = []
        failures: list[str] = []
        for path in files:
            try:
                result = run_file(path, options)
                results.append(result)
                backup = result.backup_path.name if result.backup_path else "未创建"
                self.after(0, self.write_log, f"完成：{path.name} → {result.sheet_name}；成功 {result.successful}，排除 {result.excluded}，异常 {result.errors}，基金金额 {result.total_fund:,.2f}；备份 {backup}\n")
            except Exception as exc:  # Keep batch processing independent per file.
                failures.append(f"{path.name}：{exc}")
                self.after(0, self.write_log, f"失败：{path.name} → {exc}\n")
        self.after(0, self.done, results, failures)

    def done(self, results: list[FileResult], failures: list[str]) -> None:
        self.running = False
        self.run_button.configure(state="normal")
        self.pick_button.configure(state="normal")
        self.clear_button.configure(state="normal")
        message = f"处理完成：成功 {len(results)} 个文件"
        if failures:
            message += f"，失败 {len(failures)} 个文件"
        messagebox.showinfo("完成", message)

    def on_close(self) -> None:
        if self.running:
            messagebox.showwarning("正在处理", "文件正在计算和安全保存，请等待处理完成后再关闭程序。")
            return
        self.destroy()

    def write_log(self, message: str) -> None:
        self.log.configure(state="normal")
        self.log.insert("end", message)
        self.log.see("end")
        self.log.configure(state="disabled")


def main() -> None:
    app = CalculatorApp()
    app.mainloop()


if __name__ == "__main__":
    main()
