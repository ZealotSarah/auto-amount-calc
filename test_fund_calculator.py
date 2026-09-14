import shutil
import tempfile
import unittest
import zipfile
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from fund_calculator import CalculationError, RunOptions, insurance_bucket, resolved_visit_type, run_file, validate_options

FIXTURE_DIR = Path(__file__).with_name("test_fixtures")


class CalculatorTests(unittest.TestCase):
    @staticmethod
    def make_source(path: Path, sheet_names: tuple[str, ...] = ("业务明细",)) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name in sheet_names:
            source = workbook.create_sheet(sheet_name)
            source.append(["定点编码", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量"])
            source.append(["H1", "城乡", 100, 50, 20, 21, 1])
        workbook.save(path)
        workbook.close()

    @staticmethod
    def set_formula_cached_value(path: Path, cell_reference: str, value: str) -> None:
        with zipfile.ZipFile(path) as archive:
            contents = {name: archive.read(name) for name in archive.namelist()}
        sheet_name = "xl/worksheets/sheet1.xml"
        xml = contents[sheet_name]
        marker = f'r="{cell_reference}"'.encode()
        start = xml.index(b"<c ", xml.index(marker) - 30)
        end = xml.index(b"</c>", start) + len(b"</c>")
        cell_xml = xml[start:end]
        value_start = cell_xml.index(b"<v>") + len(b"<v>")
        value_end = cell_xml.index(b"</v>", value_start)
        contents[sheet_name] = xml[:start] + cell_xml[:value_start] + value.encode() + cell_xml[value_end:] + xml[end:]
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in contents.items():
                archive.writestr(name, data)

    def test_insurance_exact_names_and_codes(self):
        for value in ("职工", "城镇职工基本医疗保险", "310", 310, 310.0):
            self.assertEqual(insurance_bucket(value), 0, value)
        for value in ("城乡", "居民", "城乡居民", "城乡居民基本医疗保险", "390", 390, 390.0, " 城乡 "):
            self.assertEqual(insurance_bucket(value), 1, value)
        for value in ("其他", None, "", "391", "非职工", "居民补充保险", "城乡其他", "职 工"):
            self.assertIsNone(insurance_bucket(value), value)

    def test_insurance_names_and_codes_are_calculated(self):
        for header in ("险种类型", "INSUTYPE"):
            with self.subTest(header=header), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "insurance.xlsx"
                workbook = Workbook()
                workbook.active.append(["定点编码", header, "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量"])
                for insurance in ("城乡", "职工", "390", 390, 390.0, "310", 310):
                    for visit in (11, 21):
                        workbook.active.append(["H1", insurance, 100, 50, 20, visit, 1])
                for insurance in ("其他", "非职工", "居民补充保险"):
                    workbook.active.append(["H1", insurance, 100, 50, 20, 21, 1])
                workbook.save(path)
                workbook.close()
                result = run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))
                self.assertEqual((result.successful, result.excluded, result.errors), (14, 3, 0))
                # 相同分组先汇总符合范围金额，再乘比例并四舍五入。
                self.assertEqual(result.total_fund, Decimal("166.22"))

    def test_group_totals_are_calculated_before_rate_and_rounding(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "grouped.xlsx"
            workbook = Workbook()
            source = workbook.active
            source.append(["定点编码", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量"])
            source.append(["H1", "城乡", 100, 0, 10, 11, 3])
            source.append(["H1", "城乡", 100, 100, 90, 11, 4])
            source.append(["H1", "城乡", 100, 50, 0.01, 21, 5])
            source.append(["H1", "城乡", 100, 50, 0.01, 21, 6])
            workbook.save(path)
            workbook.close()

            result = run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))

            # 门诊：(10 + 90) * ((0 + 100) / (100 + 100)) = 50.00。
            # 住院：(0.01 + 0.01) * 56.84% = 0.01，不能逐行取整为 0.02。
            self.assertEqual(result.total_fund, Decimal("50.01"))
            saved = load_workbook(path, data_only=True)
            output = saved["基金测算"]
            self.assertEqual(output["F8"].value, 7)
            self.assertEqual(output["G8"].value, 2)
            self.assertEqual(output["H8"].value, 50)
            self.assertEqual(output["I8"].value, 0.5)
            self.assertEqual(output["F9"].value, 11)
            self.assertEqual(output["G9"].value, 2)
            self.assertEqual(output["H9"].value, 0.01)
            self.assertEqual(next(iter(output.tables.values())).ref, "A7:I9")
            saved.close()

    def test_fixed_rate_uses_same_rate_for_inpatient_and_outpatient(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixed-rate.xlsx"
            workbook = Workbook()
            source = workbook.active
            source.append(["定点编码", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量"])
            source.append(["H1", "城乡", 0, 0, 20, 11, 2])
            source.append(["H1", "城乡", 100, 90, 20, 21, 3])
            workbook.save(path)
            workbook.close()

            result = run_file(path, RunOptions("固定比例", "自动识别", "廊坊市", "三级", None, None, True))

            self.assertEqual(result.total_fund, Decimal("22.74"))
            saved = load_workbook(path, read_only=True, data_only=True)
            try:
                output = saved["基金测算"]
                rows = {output.cell(row, 4).value: row for row in (8, 9)}
                self.assertEqual(output.cell(rows["门诊"], 8).value, 11.37)
                self.assertEqual(output.cell(rows["门诊"], 9).value, 0.5684)
                self.assertEqual(output.cell(rows["住院"], 8).value, 11.37)
                self.assertEqual(output.cell(rows["住院"], 9).value, 0.5684)
                self.assertEqual(output["K3"].value, "0.9.0")
                self.assertEqual(output["K6"].value, "普通汇总表（非 Excel 原生透视表）")
            finally:
                saved.close()

    def test_cached_formula_values_are_used_and_formulas_are_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "formula.xlsx"
            workbook = Workbook()
            source = workbook.active
            source.append(["定点编码", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量"])
            source.append(["H1", "城乡", "=50+50", 50, 20, 21, 1])
            workbook.save(path)
            workbook.close()
            self.set_formula_cached_value(path, "C2", "100")

            result = run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))

            self.assertEqual(result.total_fund, Decimal("11.37"))
            saved = load_workbook(path, data_only=False)
            try:
                self.assertEqual(saved.active["C2"].value, "=50+50")
            finally:
                saved.close()

    def test_uncached_formula_fails_with_actionable_message(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "uncached-formula.xlsx"
            workbook = Workbook()
            source = workbook.active
            source.append(["定点编码", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量"])
            source.append(["H1", "城乡", "=50+50", 50, 20, 21, 1])
            workbook.save(path)
            workbook.close()
            original_bytes = path.read_bytes()

            with self.assertRaisesRegex(CalculationError, "公式没有缓存结果"):
                run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))
            self.assertEqual(path.read_bytes(), original_bytes)

    def test_institution_name_is_first_nonempty_name_for_every_group(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "names.xlsx"
            workbook = Workbook()
            source = workbook.active
            source.append(["定点编码", "定点名称", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量"])
            source.append(["H1", None, "城乡", 100, 50, 20, 11, 1])
            source.append(["H1", "医院甲", "职工", 100, 50, 20, 21, 1])
            workbook.save(path)
            workbook.close()

            run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))

            saved = load_workbook(path, read_only=True, data_only=True)
            try:
                output = saved["基金测算"]
                self.assertEqual(output["B8"].value, "医院甲")
                self.assertEqual(output["B9"].value, "医院甲")
            finally:
                saved.close()

    def test_all_failed_rows_fail_the_file_without_modifying_it(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "all-failed.xlsx"
            workbook = Workbook()
            source = workbook.active
            source.append(["定点编码", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量"])
            source.append(["H1", "城乡", 0, 50, 20, 11, 1])
            workbook.save(path)
            workbook.close()
            original_bytes = path.read_bytes()

            with self.assertRaisesRegex(CalculationError, "没有成功计算的明细"):
                run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))
            self.assertEqual(path.read_bytes(), original_bytes)
            self.assertFalse(list(path.parent.glob("*_测算前备份_*.xlsx")))

    def test_substitution_without_violation_requires_price_column(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "missing-price.xlsx"
            self.make_source(path)
            with self.assertRaisesRegex(CalculationError, "单价.*列"):
                run_file(path, RunOptions("串换", "自动识别", "廊坊市", "三级", Decimal("1"), Decimal("1"), True))

    def test_substitution_sums_quantity_after_deduction(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "substitution.xlsx"
            workbook = Workbook()
            source = workbook.active
            source.append(["定点编码", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量", "违规金额"])
            source.append(["H1", "城乡", 100, 50, 20, 21, 5, 10])
            source.append(["H1", "城乡", 100, 50, 20, 21, 4, 20])
            workbook.save(path)
            workbook.close()

            with self.assertRaisesRegex(CalculationError, "扣减数量"):
                run_file(path, RunOptions("串换", "自动识别", "廊坊市", "三级", None, None, True))
            result = run_file(path, RunOptions("串换", "自动识别", "廊坊市", "三级", None, Decimal("2"), True))

            self.assertEqual(result.total_fund, Decimal("17.05"))
            saved = load_workbook(path, data_only=True)
            output = saved["基金测算"]
            self.assertEqual(output["F8"].value, 5)
            self.assertEqual(output["H8"].value, 17.05)
            self.assertEqual(output["K4"].value, "未填写")
            self.assertEqual(output["K5"].value, 2)
            saved.close()

    def test_code_priority_and_name_fallback(self):
        for code in ("11", 14, 51.0):
            self.assertEqual(resolved_visit_type(code, "自动识别", "普通住院"), "门诊")
        for code in ("21", 22, 52.0):
            self.assertEqual(resolved_visit_type(code, "自动识别", "普通门诊"), "住院")
        for raw, name in (("普通门诊", None), (None, "普通门诊"), ("未知", "普通门诊")):
            self.assertEqual(resolved_visit_type(raw, "自动识别", name), "门诊")
        self.assertEqual(resolved_visit_type("11", "住院"), "住院")
        with self.assertRaises(CalculationError):
            resolved_visit_type("99", "自动识别", "未知")

    def test_exclusion_totals_and_original_row_numbers(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            workbook = Workbook()
            source = workbook.active
            source.title = "明细"
            source.append(["定点编码", "定点名称", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额1", "医疗类别", "医疗类别名称", "数量"])
            source.append(["H1", "医院", "职工", 100, 50, 20, 11, "普通住院", 2])
            source.append(["H1", "医院", "城乡居民基本医疗保险", 200, 100, 30, 21, "普通门诊", 3])
            source.append(["H1", "医院", "其他", None, None, None, 11, "普通门诊", None])
            source.append(["H1", "医院", "其他", 999, 999, 999, 21, "普通住院", 9])
            source.append(["H1", "医院", None, 999, 999, 999, 11, "普通门诊", 9])
            source.append([None] * 9)
            source.append(["H1", "医院", "职工", 100, 50, 20, 99, "未知", 1])
            original = list(source.values)
            workbook.save(path)
            workbook.close()
            result = run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))
            self.assertEqual((result.processed, result.successful, result.excluded, result.errors), (6, 2, 3, 1))
            self.assertEqual(result.total_fund, Decimal("27.05"))
            self.assertTrue(result.warnings[0].startswith("第 8 行："))
            saved = load_workbook(path)
            self.assertEqual(list(saved["明细"].values), original)
            output = saved["基金测算"]
            self.assertEqual(output["F5"].value, 3)
            self.assertEqual(sum(output.cell(row, 5).value for row in (8, 9)), 300)
            self.assertEqual(output["F8"].value, 2)
            self.assertEqual(output["I8"].value, 0.5)
            self.assertEqual(output["F9"].value, 3)
            self.assertEqual(output["I9"].value, 0.5684)
            saved.close()

    def test_name_only_column(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "name.xlsx"
            workbook = Workbook()
            workbook.active.append(["定点编码", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别名称", "数量"])
            workbook.active.append(["H1", "居民", 100, 50, 20, "普通门诊", 1])
            workbook.save(path)
            workbook.close()
            result = run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))
            self.assertEqual((result.successful, result.errors, result.total_fund), (1, 0, Decimal("10.00")))

    def test_safe_save_keeps_exact_precalculation_backup(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "safe.xlsx"
            self.make_source(path)
            original_bytes = path.read_bytes()

            result = run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))

            self.assertIsNotNone(result.backup_path)
            self.assertTrue(result.backup_path.exists())
            self.assertEqual(result.backup_path.read_bytes(), original_bytes)
            self.assertNotEqual(path.read_bytes(), original_bytes)
            saved = load_workbook(path, read_only=True, data_only=True)
            self.assertIn("基金测算", saved.sheetnames)
            saved.close()

    def test_multiple_source_sheets_require_explicit_selection(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "multiple.xlsx"
            self.make_source(path, ("明细一", "明细二"))
            original_bytes = path.read_bytes()
            options = RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True)

            with self.assertRaisesRegex(CalculationError, "多个可用业务 Sheet"):
                run_file(path, options)
            self.assertEqual(path.read_bytes(), original_bytes)
            self.assertFalse(list(path.parent.glob("*_测算前备份_*.xlsx")))

            options.source_sheet = "明细二"
            result = run_file(path, options)
            self.assertEqual(result.total_fund, Decimal("11.37"))
            saved = load_workbook(path, read_only=True, data_only=True)
            try:
                self.assertEqual(saved["基金测算"]["D2"].value, "明细二")
            finally:
                saved.close()

    def test_ambiguous_candidate_columns_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ambiguous.xlsx"
            workbook = Workbook()
            workbook.active.append([
                "定点编码", "险种类型", "医疗费总额", "基金支付总额",
                "符合范围金额1", "符合范围金额2", "医疗类别", "数量",
            ])
            workbook.active.append(["H1", "城乡", 100, 50, 20, 30, 21, 1])
            workbook.save(path)
            workbook.close()

            with self.assertRaisesRegex(CalculationError, "符合范围金额匹配到多个列"):
                run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))

    def test_required_choices_and_deductions_are_validated(self):
        valid = RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True)
        invalid_options = (
            (RunOptions("通用", "自动识别", "", "三级", None, None, True), "参保地"),
            (RunOptions("通用", "自动识别", "廊坊市", "", None, None, True), "医疗机构级别"),
            (RunOptions("通用", "自动识别", "廊坊市", "三级", Decimal("-1"), None, True), "扣减单价"),
            (RunOptions("串换", "自动识别", "廊坊市", "三级", None, Decimal("-1"), True), "扣减数量"),
            (RunOptions("串换", "自动识别", "廊坊市", "三级", None, Decimal("NaN"), True), "扣减数量"),
        )
        validate_options(valid)
        for options, message in invalid_options:
            with self.subTest(message=message), self.assertRaisesRegex(CalculationError, message):
                validate_options(options)

    def test_sanitized_sample_02_regression(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample_02_sanitized.xlsx"
            shutil.copy2(FIXTURE_DIR / path.name, path)

            result = run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True, "With"))

            self.assertEqual((result.processed, result.successful, result.excluded, result.errors), (6, 5, 1, 0))
            self.assertEqual(result.total_fund, Decimal("153.13"))
            saved = load_workbook(path, read_only=True, data_only=True)
            try:
                output = saved["基金测算"]
                self.assertEqual(output["D2"].value, "With")
                self.assertEqual(sum(output.cell(row, 6).value for row in range(8, 12)), 12)
            finally:
                saved.close()

    def test_sanitized_sample_03_regression(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample_03_sanitized.xlsx"
            shutil.copy2(FIXTURE_DIR / path.name, path)

            result = run_file(path, RunOptions("固定比例", "自动识别", "廊坊市", "三级", None, None, True, "With"))

            self.assertEqual((result.processed, result.successful, result.excluded, result.errors), (4, 3, 1, 0))
            self.assertEqual(result.total_fund, Decimal("102.04"))
            saved = load_workbook(path, read_only=True, data_only=True)
            try:
                output = saved["基金测算"]
                self.assertEqual({output.cell(row, 2).value for row in range(8, 11)}, {"脱敏医院乙"})
                self.assertEqual({output.cell(row, 9).value for row in range(8, 10)}, {0.5684})
            finally:
                saved.close()

    def test_visit_count_does_not_dedupe_same_person(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "dedupe.xlsx"
            workbook = Workbook()
            source = workbook.active
            source.append(["定点编码", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量"])
            # 同一人员（相同编码/身份）出现 3 条明细，人次应计 3，不去重。
            source.append(["H1", "城乡", 100, 50, 20, 21, 1])
            source.append(["H1", "城乡", 100, 50, 20, 21, 1])
            source.append(["H1", "城乡", 100, 50, 20, 21, 1])
            workbook.save(path)
            workbook.close()

            result = run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))

            self.assertEqual(result.successful, 3)
            saved = load_workbook(path, data_only=True)
            output = saved["基金测算"]
            self.assertEqual(output["G8"].value, 3)
            saved.close()

    def test_visit_count_unchanged_by_substitution_deduction(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "subst-count.xlsx"
            workbook = Workbook()
            source = workbook.active
            source.append(["定点编码", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量", "违规金额"])
            source.append(["H1", "城乡", 100, 50, 20, 21, 5, 10])
            source.append(["H1", "城乡", 100, 50, 20, 21, 4, 20])
            workbook.save(path)
            workbook.close()

            result = run_file(path, RunOptions("串换", "自动识别", "廊坊市", "三级", None, Decimal("2"), True))

            self.assertEqual(result.successful, 2)
            saved = load_workbook(path, data_only=True)
            output = saved["基金测算"]
            self.assertEqual(output["G8"].value, 2)
            self.assertEqual(output["F8"].value, 5)
            saved.close()

    def test_visit_count_sum_equals_successful(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "count-sum.xlsx"
            workbook = Workbook()
            source = workbook.active
            source.append(["定点编码", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量"])
            source.append(["H1", "城乡", 100, 50, 20, 21, 1])
            source.append(["H1", "城乡", 100, 50, 20, 11, 1])
            source.append(["H2", "职工", 100, 50, 20, 21, 1])
            workbook.save(path)
            workbook.close()

            result = run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))

            self.assertEqual(result.successful, 3)
            saved = load_workbook(path, data_only=True)
            output = saved["基金测算"]
            count_total = sum(output.cell(row, 7).value for row in range(8, 11))
            self.assertEqual(count_total, result.successful)
            saved.close()

    def test_nine_column_header_and_audit_position(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nine-col.xlsx"
            workbook = Workbook()
            source = workbook.active
            source.append(["定点编码", "险种类型", "医疗费总额", "基金支付总额", "符合范围金额", "医疗类别", "数量"])
            source.append(["H1", "城乡", 100, 50, 20, 21, 1])
            workbook.save(path)
            workbook.close()

            run_file(path, RunOptions("通用", "自动识别", "廊坊市", "三级", None, None, True))

            saved = load_workbook(path, data_only=True)
            output = saved["基金测算"]
            self.assertEqual(output["A1"].value, "基金金额测算结果")
            self.assertEqual(
                [output.cell(7, column).value for column in range(1, 10)],
                ["医疗机构编码", "医疗机构名称", "险种类别", "医疗类别", "医疗总额", "数量总和", "人次", "基金金额", "报销比例"],
            )
            self.assertEqual(next(iter(output.tables.values())).ref, "A7:I8")
            self.assertEqual(output["G8"].value, 1)
            self.assertEqual(output["J2"].value, "生成时间")
            self.assertEqual(output["K3"].value, "0.9.0")
            saved.close()


if __name__ == "__main__":
    unittest.main()
